from __future__ import annotations

from io import BytesIO
from unittest import TestCase

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from RAG import (
    ApproveVersionCommand,
    FileParsingError,
    InMemoryDocumentStore,
    IngestFileDocumentCommand,
    MemoryModule,
    PublishVersionCommand,
    SearchPublishedMemoryQuery,
    UnsupportedFileTypeError,
)


class FileIngestionTestCase(TestCase):
    def test_all_supported_formats_use_text_ingestion_pipeline(self) -> None:
        fixtures = {
            "policy.txt": "Text policy requires respectful communication.".encode(),
            "policy.md": b"# Markdown policy\n\nPhishing is prohibited.",
            "policy.html": (
                b"<html><head><style>hidden css</style></head><body>"
                b"<h1>HTML policy</h1><p>Protect personal data.</p>"
                b"<script>hidden script</script></body></html>"
            ),
            "policy.docx": _docx_bytes(),
            "policy.pdf": _pdf_bytes("PDF policy prohibits credential theft."),
        }
        source_types = {
            "policy.txt": "text",
            "policy.md": "md",
            "policy.html": "html",
            "policy.docx": "docx",
            "policy.pdf": "pdf",
        }
        expected_text = {
            "policy.txt": "communication",
            "policy.md": "Phishing",
            "policy.html": "personal data",
            "policy.docx": "Retention",
            "policy.pdf": "theft",
        }

        for index, (file_name, content) in enumerate(fixtures.items()):
            with self.subTest(file_name=file_name):
                memory = MemoryModule(
                    InMemoryDocumentStore(),
                    embedding_provider=_RecordingEmbeddingProvider(),
                )
                result = memory.ingest_file_document(
                    IngestFileDocumentCommand(
                        document_id=f"doc-{index}",
                        version_id=f"version-{index}",
                        tenant_id="tenant-1",
                        title=file_name,
                        file_name=f"uploads\\{file_name}",
                        file_content=content,
                        max_chunk_chars=30,
                        chunk_overlap_chars=5,
                    )
                )

                self.assertEqual(result.version.review_status, "processed")
                self.assertEqual(result.version.source_type, source_types[file_name])
                self.assertEqual(result.version.original_name, file_name)
                self.assertTrue(result.chunks)
                self.assertTrue(all(chunk.text for chunk in result.chunks))
                self.assertTrue(all(chunk.is_embedded for chunk in result.chunks))
                self.assertIn(
                    expected_text[file_name],
                    " ".join(chunk.text for chunk in result.chunks),
                )

        html_text = " ".join(
            chunk.text
            for chunk in memory.ingest_file_document(
                IngestFileDocumentCommand(
                    document_id="html-visible",
                    version_id="html-visible-v1",
                    tenant_id="tenant-1",
                    title="HTML",
                    file_name="visible.HTML",
                    file_content=fixtures["policy.html"],
                )
            ).chunks
        )
        self.assertIn("Protect personal data", html_text)
        self.assertNotIn("hidden css", html_text)
        self.assertNotIn("hidden script", html_text)

    def test_file_can_be_moderated_published_and_searched(self) -> None:
        memory = MemoryModule(InMemoryDocumentStore())
        result = memory.ingest_file_document(
            IngestFileDocumentCommand(
                document_id="terms",
                version_id="terms-v1",
                tenant_id="tenant-1",
                title="Terms of service",
                file_name="terms.txt",
                file_content=b"Automated scraping and credential theft are prohibited.",
                trust_level="trusted",
            )
        )
        memory.approve_version(
            ApproveVersionCommand("terms", result.version.version_id, "tenant-1")
        )
        memory.publish_version(
            PublishVersionCommand("terms", result.version.version_id, "tenant-1")
        )

        matches = memory.search_published_memory(
            SearchPublishedMemoryQuery(
                tenant_id="tenant-1",
                query="credential theft",
                retrieval_mode="lexical",
            )
        )

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0].document_id, "terms")

    def test_invalid_files_are_rejected_before_document_is_saved(self) -> None:
        cases = (
            ("policy.csv", b"unsupported", UnsupportedFileTypeError),
            ("policy.txt", b"", FileParsingError),
            ("policy.txt", b"\xff\xfeinvalid", FileParsingError),
            ("policy.docx", b"not a docx", FileParsingError),
            ("policy.pdf", _pdf_bytes(None), FileParsingError),
            ("policy.html", b"<script>only hidden text</script>", FileParsingError),
        )

        for index, (file_name, content, error_type) in enumerate(cases):
            with self.subTest(file_name=file_name, index=index):
                memory = MemoryModule(InMemoryDocumentStore())
                with self.assertRaises(error_type):
                    memory.ingest_file_document(
                        IngestFileDocumentCommand(
                            document_id=f"invalid-{index}",
                            version_id=f"invalid-{index}-v1",
                            tenant_id="tenant-1",
                            title="Invalid file",
                            file_name=file_name,
                            file_content=content,
                        )
                    )


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("DOCX policy requires consent.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Retention"
    table.cell(0, 1).text = "Thirty days"
    document.save(output)
    return output.getvalue()


def _pdf_bytes(text: str | None) -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text is not None:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


class _RecordingEmbeddingProvider:
    model = "test-embedding"
    dimensions = 1536

    def embed_documents(self, texts: tuple[str, ...]) -> tuple[tuple[float, ...], ...]:
        vector = (1.0,) + (0.0,) * (self.dimensions - 1)
        return tuple(vector for _ in texts)

    def embed_query(self, text: str) -> tuple[float, ...]:
        del text
        return (1.0,) + (0.0,) * (self.dimensions - 1)
